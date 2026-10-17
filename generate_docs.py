import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()
    
    # Page setup - Margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(10.5)
    normal_font.color.rgb = RGBColor(33, 37, 41)
    
    # Title Banner
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("VILLAGEPULSE (GRAMPULSE)")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 120, 87) # Emerald Green
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Comprehensive Project Master Documentation & AI Context Manual")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    # Metadata Box
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    metadata = [
        ("Project Name", "VillagePulse (Smart Rural Survey & Family Intelligence System)"),
        ("Live Deployment URL", "https://villagepulse.netlify.app (Netlify SPA with PWA Cache)"),
        ("Repository", "https://github.com/nareshkumarregar89-bit/GramPulse.git (Branch: main)"),
        ("Documentation Purpose", "Instant full-context onboarding for Future AI Assistants & Developers")
    ]
    
    for i, (k, v) in enumerate(metadata):
        c0 = meta_table.cell(i, 0)
        c1 = meta_table.cell(i, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.8)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 80, 80, 100, 100)
        set_cell_margins(c1, 80, 80, 100, 100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(51, 65, 85)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(15, 23, 42)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(13, 94, 69)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)
        return h

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r0 = p.add_run(bold_prefix + ": ")
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(30, 41, 59)
        r1 = p.add_run(text)
        r1.font.color.rgb = RGBColor(51, 65, 85)

    def add_callout(text, bg="EFF6FF", border="3B82F6"):
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0)
        c.width = Inches(7.0)
        set_cell_background(c, bg)
        set_cell_margins(c, 100, 100, 140, 140)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 58, 138)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # SECTION 1: EXECUTIVE SUMMARY
    add_h1("1. Executive Summary & Purpose")
    p = doc.add_paragraph(
        "VillagePulse (formerly GramPulse) is a modern, cloud-synchronized, Progressive Web Application (PWA) "
        "designed specifically for Indian rural administrative governance, village surveyors, and panchayat representatives. "
        "The system digitizes household door-to-door demographic surveys, family genealogies, socioeconomic status, "
        "government welfare IDs (Jan Aadhaar, Ration Card), and provides real-time community intelligence, multi-tier RBAC "
        "(Role-Based Access Control), dynamic village isolation, and official bilingual A4 dossier printing."
    )
    p.paragraph_format.space_after = Pt(6)
    
    add_callout(
        "CRITICAL NOTICE FOR FUTURE AI ASSISTANTS: This codebase connects directly to a live production Firebase Firestore "
        "database. Mock and dummy data generators have been completely purged. Real verified records for villages like Samodi "
        "and Kotdi are active in production. Super Admin UID: sJwiZZ6cf8Ze8YAmMMK8T1FwF232."
    )

    # SECTION 2: TECHNOLOGY STACK & ARCHITECTURE
    add_h1("2. Technology Stack & Directory Architecture")
    add_bullet("Core Frontend", "React 18.3.1 with Vite 5.4.11 bundling for instant HMR and optimized production chunks.")
    add_bullet("Styling Engine", "Tailwind CSS 3.4.17 with full Dark Mode support synced to document.documentElement.classList.")
    add_bullet("Cloud Backend", "Firebase v13 Cloud Firestore (Real-time NoSQL) and Firebase Authentication.")
    add_bullet("State Management", "React Context API (AuthContext, DataContext, ToastContext) paired with useLocalStorage fallback cache.")
    add_bullet("Icons & Assets", "Lucide React (over 40+ icons used across dashboards, cards, and modal dialogs).")
    add_bullet("Printing & Reporting", "Custom CSS @media print styling with React Portal rendering directly into document.body, jsPDF 4.2.1, and XLSX (SheetJS) 0.18.5.")
    add_bullet("PWA & Offline", "Vite PWA plugin (v2.0.0) with custom Service Worker, auto-update prompt, and Netlify cache headers.")

    add_h2("Repository Directory Structure")
    structure_text = (
        "GramPulse/\n"
        "├── public/                  # Favicons, logo.png, PWA icons, manifest\n"
        "├── src/\n"
        "│   ├── components/          # Reusable UI components\n"
        "│   │   ├── ExcelImportModal.jsx   # Bulk family import with Excel parsing\n"
        "│   │   ├── PrintDataModal.jsx     # Production A4 print engine with React Portal\n"
        "│   │   ├── TiltCard.jsx           # Premium 3D interactive physics tilt card\n"
        "│   │   ├── UpdatePrompt.jsx       # Service worker PWA update notifications\n"
        "│   │   └── UserProfileModal.jsx   # Profile editor & password management\n"
        "│   ├── contexts/            # Global application state\n"
        "│   │   ├── AuthContext.jsx        # RBAC, login/register, Google auth, Firestore role sync\n"
        "│   │   ├── DataContext.jsx        # Families, villages, mohallas, castes, sync listeners\n"
        "│   │   └── ToastContext.jsx       # Global notification toasts\n"
        "│   ├── data/                # Initial configurations & production seeds\n"
        "│   │   ├── castes.js              # Supported caste categories\n"
        "│   │   ├── mohallas.js            # Default approved mohalla list\n"
        "│   │   ├── villageData.js         # Production Samodi initial family seeds\n"
        "│   │   └── villages.js            # Initial village records (Samodi, Kotdi)\n"
        "│   ├── layouts/             # Responsive app shell\n"
        "│   │   ├── BottomNavigation.jsx   # Modern floating capsule mobile taskbar\n"
        "│   │   ├── Layout.jsx             # Main shell wrapper with desktop sidebar\n"
        "│   │   ├── Sidebar.jsx            # Desktop navigation sidebar with collapse\n"
        "│   │   └── TopNavbar.jsx          # Top bar with village selector, lang, darkmode, profile\n"
        "│   ├── pages/               # Primary application routes\n"
        "│   │   ├── AdminPanel.jsx         # User management, village allocation, wiping\n"
        "│   │   ├── Castes.jsx             # Caste distribution analytics\n"
        "│   │   ├── Dashboard.jsx          # KPI metrics, isometric village hero, charts\n"
        "│   │   ├── DeletedRecords.jsx     # Recycle bin & restore system\n"
        "│   │   ├── Families.jsx           # Paginated family directory with filters & print\n"
        "│   │   ├── FamilyProfile.jsx      # Individual family dossier & government IDs\n"
        "│   │   ├── Login.jsx              # Dual authentication (Email/Password & Google)\n"
        "│   │   ├── Mohallas.jsx           # Mohalla demographics\n"
        "│   │   ├── NewSurvey.jsx          # Multi-step household survey wizard with GPS\n"
        "│   │   ├── Reports.jsx            # Statistical analytical reports & export\n"
        "│   │   ├── Search.jsx             # Advanced multi-parameter fuzzy search\n"
        "│   │   ├── Settings.jsx           # App configurations, config sync, cache purge\n"
        "│   │   └── VillageManagement.jsx  # Dynamic village CRUD & status toggling\n"
        "│   ├── firebase.js          # Initialized Firebase app, db, auth, provider\n"
        "│   ├── index.css            # Tailored styling, animations, and @media print rules\n"
        "│   └── main.jsx             # React DOM root entry\n"
        "├── netlify.toml             # Netlify deployment & SPA routing configuration\n"
        "└── package.json             # Project dependencies and npm scripts"
    )
    p_struct = doc.add_paragraph()
    r_struct = p_struct.add_run(structure_text)
    r_struct.font.name = 'Consolas'
    r_struct.font.size = Pt(8.5)
    r_struct.font.color.rgb = RGBColor(30, 41, 59)
    p_struct.paragraph_format.space_before = Pt(4)
    p_struct.paragraph_format.space_after = Pt(10)

    # SECTION 3: DATA MODELS & SCHEMAS
    add_h1("3. Core Schemas & Database Models")
    add_bullet("families Collection", "Primary collection storing household survey records. Document ID format: FAM-XXXX or UUID.")
    
    schema_table = doc.add_table(rows=12, cols=3)
    schema_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [("Field Name", Inches(1.8)), ("Type", Inches(1.2)), ("Description", Inches(4.0))]
    for col_idx, (h_text, width) in enumerate(headers):
        cell = schema_table.cell(0, col_idx)
        cell.width = width
        set_cell_background(cell, "107857")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
        
    fields = [
        ("id", "string", "Unique family identifier (e.g., 'FAM-0001')"),
        ("village / villageId", "string", "Village name string & canonical id ('samodi', 'kotdi')"),
        ("familyHeadName", "string", "Bilingual combined name (e.g. 'Mohan Lal Regar (मोहन लाल रेगर)')"),
        ("familyHeadNameEnglish", "string", "English name of the family head"),
        ("familyHeadNameHindi", "string", "Hindi script representation of family head"),
        ("fatherName / fatherNameHindi", "string", "Father/Husband name in English and Hindi"),
        ("mohalla", "string", "Neighborhood colony (e.g. 'Regar Mohalla', 'Ahir Mohalla')"),
        ("caste", "string", "Caste grouping (e.g. 'Regar', 'Ahir', 'Bheel', 'Gurjar')"),
        ("janAadhaar / rationCard", "string", "Government entitlement ID numbers"),
        ("members", "array of objects", "Array of individual family members: name, relation, age, gender, dob, aadhaar, education"),
        ("surveyorId / surveyorName", "string", "UID and name of the field surveyor who recorded the family")
    ]
    
    for row_idx, (f_name, f_type, f_desc) in enumerate(fields, start=1):
        bg = "FFFFFF" if row_idx % 2 == 1 else "F8FAFC"
        c0 = schema_table.cell(row_idx, 0)
        c1 = schema_table.cell(row_idx, 1)
        c2 = schema_table.cell(row_idx, 2)
        c0.width = Inches(1.8)
        c1.width = Inches(1.2)
        c2.width = Inches(4.0)
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_background(c2, bg)
        set_cell_margins(c0, 60, 60, 80, 80)
        set_cell_margins(c1, 60, 60, 80, 80)
        set_cell_margins(c2, 60, 60, 80, 80)
        
        c0.paragraphs[0].add_run(f_name).font.size = Pt(9.0)
        c1.paragraphs[0].add_run(f_type).font.size = Pt(9.0)
        c2.paragraphs[0].add_run(f_desc).font.size = Pt(9.0)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # SECTION 4: SECURITY & RBAC
    add_h1("4. Security Architecture & Role-Based Access Control (RBAC)")
    add_bullet("Super Admin Role", "Users matching ADMIN_UID ('sJwiZZ6cf8Ze8YAmMMK8T1FwF232') or having role: 'admin' in Firestore users/{uid}. They have unrestricted read/write access to all villages, user management, and database purging.")
    add_bullet("Surveyor Role", "Field surveyors who can only record new families and view/edit households assigned to their specific village(s) or created under their surveyorId.")
    add_bullet("Unassigned User Workflow", "When a new user signs up, their profile defaults to unassigned with no village rights until the Super Admin activates and assigns them a village in the Admin Panel.")
    add_bullet("Real-Time Role Sync", "AuthContext listens to Firestore users/{uid} using onSnapshot. Any change in status, role, or assigned villages takes effect instantly on the client without requiring logout or page reload.")

    # SECTION 5: PRINT ENGINE ARCHITECTURE
    add_h1("5. The Official Print Engine Architecture")
    p_print = doc.add_paragraph(
        "A major challenge encountered in web applications during window.print() is that background UI components "
        "(such as dashboard charts, headers, and sidebars) bleed into the print preview. "
        "In VillagePulse, we solved this with an isolated React Portal architecture:"
    )
    p_print.paragraph_format.space_after = Pt(4)
    add_bullet("Direct Body Portalling", "PrintDataModal renders its internal <PrintDocument> directly into document.body using ReactDOM.createPortal(<PrintDocument />, document.body) under the CSS class .print-portal-area.")
    add_bullet("Dynamic CSS Print Rules", "index.css enforces body.is-printing-report #root { display: none !important; }. When printing starts, the entire React application root is concealed from the browser print spooler.")
    add_bullet("A4 Page Pagination", "Tables utilize page-break-inside: avoid, official emblem headers, bilingual column captions, and print-specific footer counters.")

    # SECTION 6: COMPLETE GIT COMMIT HISTORY
    add_h1("6. Chronological Git Commit History & Release Milestones")
    
    commits = [
        ("ef7b304", "fix(print): portal print document to body and hide root; remove dummy data; modernize mobile bottom bar", "Current Release: Fixed printing bug, completely eliminated mock generators and UI toggles, redesigned floating capsule mobile taskbar."),
        ("a8b2c6d", "feat: add on-demand dummy data toggle feature for demo preview", "Temporary demo toggle feature (now removed in production)."),
        ("30fa1a1", "fix: import Clock icon from lucide-react in AdminPanel", "Fixed unimported Clock icon reference causing runtime error in admin panel."),
        ("8f6d17b", "feat: eliminate mock data auto-seeding and add clearAllSurveyData", "Removed automatic local storage mock seeding and added batch deletion for clean slate."),
        ("470e912", "fix: ensure update button immediately reloads with controllerchange", "PWA service worker instant reload enhancement."),
        ("4bb580c", "style: hide netlify drawer and feedback badge in global css", "Cleaned Netlify preview feedback widgets from UI."),
        ("e2daf68", "chore: rename app from GramPulse to VillagePulse across all files", "Official rebranding to VillagePulse with logo and bilingual metadata."),
        ("f691993", "feat: add premium responsive top-bar refresh button", "Added animated sync and refresh button for live Firestore updates."),
        ("d4220fd", "feat: enable realtime role and profile sync without page refresh", "Implemented onSnapshot on users/{uid} in AuthContext."),
        ("09c44ee", "feat: add pwa update prompt with netlify cache headers", "Added PWA update prompt toast and Netlify cache headers."),
        ("103f07d", "feat: complete Firebase Auth user deletion and unassigned workflow", "Full surveyor activation and deletion workflow in AdminPanel."),
        ("ba79b3a", "feat: implement premium 3D UI, isometric village hero, tilt cards", "Modern dashboard aesthetic with glassmorphism and 3D cards."),
        ("0d1f098", "feat: set default mohallas to only Regar, Ahir, Bheel, Gurjar", "Localized mohallas to target village wards."),
        ("150c4e3", "feat: dynamic village management with Samodi and Kotdi", "Multi-village support with village isolation and switcher."),
        ("0499f2e", "Add Google Sign-In authentication support with Firebase", "Added Google popup authentication."),
        ("6361fb0", "Add Professional Data Printing Feature with A4 layout", "Initial implementation of the multi-family print engine."),
        ("ed8360e", "Add Firebase Authentication with roles, Surveyor tracking, GPS", "Authentication foundation and geolocation logging."),
        ("bbec6d9", "Configure PWA with VitePWA plugin, manifest, service worker", "Progressive Web App installation capabilities."),
        ("d224b37", "Integrate Firebase Firestore for global real-time synchronization", "Transitioned from local-only storage to live Firestore.")
    ]
    
    commit_table = doc.add_table(rows=len(commits) + 1, cols=3)
    commit_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_headers = [("Hash", Inches(1.1)), ("Commit Subject", Inches(2.7)), ("Significance & Details", Inches(3.2))]
    for col_idx, (h_text, width) in enumerate(c_headers):
        cell = commit_table.cell(0, col_idx)
        cell.width = width
        set_cell_background(cell, "334155")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.0)

    for idx, (chash, csubj, cdesc) in enumerate(commits, start=1):
        bg = "FFFFFF" if idx % 2 == 1 else "F8FAFC"
        c0 = commit_table.cell(idx, 0)
        c1 = commit_table.cell(idx, 1)
        c2 = commit_table.cell(idx, 2)
        c0.width = Inches(1.1)
        c1.width = Inches(2.7)
        c2.width = Inches(3.2)
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_background(c2, bg)
        set_cell_margins(c0, 60, 60, 80, 80)
        set_cell_margins(c1, 60, 60, 80, 80)
        set_cell_margins(c2, 60, 60, 80, 80)
        
        r0 = c0.paragraphs[0].add_run(chash)
        r0.font.name = 'Consolas'
        r0.font.size = Pt(8.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(16, 120, 87)
        
        r1 = c1.paragraphs[0].add_run(csubj)
        r1.font.size = Pt(8.5)
        r1.font.bold = True
        
        r2 = c2.paragraphs[0].add_run(cdesc)
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # SECTION 7: DEVELOPER & AI GUIDE
    add_h1("7. Rules of Engagement for Future AI Coding Assistants")
    add_bullet("1. DO NOT RE-INTRODUCE DUMMY DATA", "The client has explicitly requested a clean, real production database. Never add mock generator loops, sample data seeders, or fake family generators.")
    add_bullet("2. RESPECT DESKTOP & MOBILE SEPARATION", "The desktop/laptop layout is verified as rock-solid. Mobile navigation utilizes the floating capsule dock in src/layouts/BottomNavigation.jsx.")
    add_bullet("3. RESPECT BILINGUAL STANDARDS", "All citizen and caste records must support English and Hindi translations (e.g. familyHeadNameEnglish, familyHeadNameHindi).")
    add_bullet("4. MAINTAIN PRINT PORTAL INTEGRITY", "Whenever modifying printing, ensure <PrintDocument> is rendered via React Portal directly into document.body to avoid parent CSS clipping or dashboard leakage.")
    add_bullet("5. LOCALSTORAGE + FIRESTORE FALLBACK", "The app uses a hybrid caching strategy where onSnapshot synchronizes to React state and useLocalStorage ensures smooth offline loading.")

    # Save Document
    file_path = "PROJECT_MASTER_DOCUMENTATION.docx"
    doc.save(file_path)
    print(f"Successfully generated {file_path}")

if __name__ == "__main__":
    create_document()
